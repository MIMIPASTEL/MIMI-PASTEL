"""Analisis detallado de la hoja de produccion."""
import sys
from collections import Counter, defaultdict
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=False)
ws = wb["prod1"]

print("== ENCABEZADOS COMPLETOS (fila 1) ==")
for c in range(1, ws.max_column + 1):
    v = ws.cell(row=1, column=c).value
    if v is not None:
        print(f"  col {c:2d} ({ws.cell(row=1, column=c).coordinate:>4}): {v!r}")

print("\n== FILA 2 completa ==")
for c in range(1, ws.max_column + 1):
    v = ws.cell(row=2, column=c).value
    if v is not None:
        print(f"  {ws.cell(row=2, column=c).coordinate}: {v!r}")

print("\n== VALORES DISTINTOS COLUMNA B (clientes/canales) ==")
clientes = Counter()
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=2).value
    if v is not None:
        clientes[str(v).strip()] += 1
for k, n in clientes.most_common():
    print(f"  {k!r}: {n}")

print("\n== FECHAS ==")
fechas = []
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if v is not None:
        fechas.append((r, v))
print(f"  filas con fecha: {len(fechas)}")
print(f"  primera: r{fechas[0][0]} = {fechas[0][1]!r}")
print(f"  ultima : r{fechas[-1][0]} = {fechas[-1][1]!r}")
tipos = Counter(type(v).__name__ for _, v in fechas)
print(f"  tipos: {dict(tipos)}")
print("  ultimas 6 fechas:", [v for _, v in fechas[-6:]])

print("\n== FILAS 'TOTAL' ==")
tot = [r for r in range(2, ws.max_row + 1) if str(ws.cell(row=r, column=2).value).strip().upper() == "TOTAL"]
print(f"  cantidad: {len(tot)} | primeras: {tot[:5]} | ultimas: {tot[-3:]}")

print("\n== ULTIMAS 12 FILAS CON DATO ==")
for r in range(max(2, ws.max_row - 12), ws.max_row + 1):
    vals = [f"{ws.cell(row=r,column=c).coordinate}={ws.cell(row=r,column=c).value!r}"
            for c in range(1, ws.max_column + 1) if ws.cell(row=r, column=c).value is not None]
    if vals:
        print(f"  r{r}: " + " | ".join(vals[:8]))

print("\n== FORMULAS DISTINTAS ==")
formulas = Counter()
for row in ws.iter_rows(min_row=2):
    for cell in row:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            formulas[cell.value] += 1
for f, n in formulas.most_common(10):
    print(f"  {f}  x{n}")
print(f"  total celdas con formula: {sum(formulas.values())}")

print("\n== ANCHOS / ESTILOS BASICOS ==")
print("  anchos:", {k: round(v.width, 1) for k, v in list(ws.column_dimensions.items())[:12] if v.width})
print("  filas congeladas:", ws.freeze_panes)
