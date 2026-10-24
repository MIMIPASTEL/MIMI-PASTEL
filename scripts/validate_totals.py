"""Valida los totales por dia del libro contra la suma real de sus filas."""
import sys
from datetime import date, timedelta
from openpyxl import load_workbook

path = sys.argv[1]
wbf = load_workbook(path, data_only=False)   # formulas
wbv = load_workbook(path, data_only=True)    # valores cacheados por Excel

wsf = wbf["prod1"]
wsv = wbv["prod1"]

COLS = list(range(3, 21))  # C..T

def num(v):
    return v if isinstance(v, (int, float)) else 0

# localizar bloques: cada TOTAL cierra el bloque desde el TOTAL anterior
totales = [r for r in range(2, wsf.max_row + 1)
           if str(wsf.cell(row=r, column=2).value).strip().upper() == "TOTAL"]

problemas = []
sin_cache = 0

inicio = 2
for t in totales:
    detalle = range(inicio, t)
    for c in COLS + [21]:  # productos + columna U (TOT)
        esperado = sum(num(wsv.cell(row=r, column=c).value) for r in detalle)
        cache = wsv.cell(row=t, column=c).value
        if cache is None:
            sin_cache += 1
            continue
        if abs(float(cache) - float(esperado)) > 0.001:
            problemas.append((t, wsf.cell(row=t, column=c).coordinate,
                              float(cache), float(esperado), wsv.cell(row=t, column=1).value))
    inicio = t + 1

print(f"Bloques/dias cerrados con TOTAL: {len(totales)}")
print(f"Celdas sin valor cacheado: {sin_cache}")
print(f"Discrepancias: {len(problemas)}")
for p in problemas[:25]:
    print(f"  fila {p[0]:3d} {p[1]:>4}  en hoja={p[2]:>8.1f}  real={p[3]:>8.1f}  dif={p[2]-p[3]:>7.1f}  fecha={p[4]}")

# filas que quedan fuera de todo bloque
if inicio <= wsf.max_row:
    print(f"\nFilas despues del ultimo TOTAL: {inicio}..{wsf.max_row}")
