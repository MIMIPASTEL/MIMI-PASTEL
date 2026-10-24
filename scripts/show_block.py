"""Muestra un bloque completo para entender la aritmetica real."""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
ini, fin = int(sys.argv[2]), int(sys.argv[3])

wbf = load_workbook(path, data_only=False)
wsv = load_workbook(path, data_only=True)
wsf, wsv_s = wbf["prod1"], wsv["prod1"]

for r in range(ini, fin + 1):
    a = wsf.cell(row=r, column=1).value
    b = wsf.cell(row=r, column=2).value
    fc = wsf.cell(row=r, column=3).value
    vc = wsv_s.cell(row=r, column=3).value
    fu = wsf.cell(row=r, column=21).value
    vu = wsv_s.cell(row=r, column=21).value
    print(f"r{r:3d} A={a!r:>10} B={b!r:<16} C: formula={fc!r:<18} valor={vc!r:<6} | U: {fu!r:<20} = {vu!r}")
