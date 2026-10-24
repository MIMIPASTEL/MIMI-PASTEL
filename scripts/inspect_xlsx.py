"""Inspeccion estructural de un libro Excel (solo lectura)."""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=False)
print("Archivo:", path)
print("Hojas:", wb.sheetnames)
print()

for ws in wb.worksheets:
    print("=" * 70)
    print(f"HOJA: {ws.title}  |  dims={ws.dimensions}  |  max_row={ws.max_row}  max_col={ws.max_column}")
    print(f"  estado={ws.sheet_state}")
    if ws.merged_cells.ranges:
        print(f"  combinadas: {[str(r) for r in ws.merged_cells.ranges][:10]}")
    # formulas y valores en las primeras filas
    for r in range(1, min(ws.max_row, 25) + 1):
        vals = []
        for c in range(1, min(ws.max_column, 14) + 1):
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if v is None:
                continue
            if isinstance(v, str) and len(v) > 40:
                v = v[:40] + "…"
            vals.append(f"{cell.coordinate}={v!r}")
        if vals:
            print(f"  r{r}: " + " | ".join(vals))
    print()
