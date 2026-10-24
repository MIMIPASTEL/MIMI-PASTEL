"""Agregados reales del libro de produccion."""
import sys
from collections import Counter, defaultdict
from datetime import date, timedelta
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path, data_only=False)
ws = wb["prod1"]

PRODUCTOS = {}
for c in range(3, 21):
    PRODUCTOS[c] = ws.cell(row=1, column=c).value

def a_fecha(v):
    if isinstance(v, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(v))
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d/%m/%y", "%d/%m/%Y"):
            try:
                d, m, y = s.split("/")
                y = int(y) + (2000 if len(y) == 2 else 0)
                return date(y, int(m), int(d))
            except Exception:
                pass
    return None

por_producto = Counter()
por_cliente = Counter()
dias = set()
filas_dato = 0
sob_venta = []

for r in range(2, ws.max_row + 1):
    b = ws.cell(row=r, column=2).value
    if b is None:
        continue
    b = str(b).strip()
    if b.upper() == "TOTAL":
        continue
    f = a_fecha(ws.cell(row=r, column=1).value)
    if f:
        dias.add(f)
    filas_dato += 1
    if b.lower() in ("sob", "venta"):
        sob_venta.append(r)
        continue
    por_cliente[b.upper()] += 1
    for c, nom in PRODUCTOS.items():
        v = ws.cell(row=r, column=c).value
        if isinstance(v, (int, float)):
            por_producto[str(nom)] += v

print(f"Dias distintos: {len(dias)}")
if dias:
    print(f"Rango: {min(dias)} -> {max(dias)}")
print(f"Filas de datos: {filas_dato}")
print(f"Filas marcadas sob/venta: {len(sob_venta)} -> {sob_venta[:10]}")

print("\n== UNIDADES POR PRODUCTO ==")
for k, v in por_producto.most_common():
    print(f"  {k:<8} {v:,.0f}")

print("\n== FILAS POR CLIENTE/CANAL (normalizado) ==")
for k, v in por_cliente.most_common(20):
    print(f"  {k:<20} {v}")

print("\n== MUESTRA FILAS 'sob' y 'venta' ==")
for r in sob_venta[:6]:
    vals = [f"{ws.cell(row=r,column=c).coordinate}={ws.cell(row=r,column=c).value!r}"
            for c in range(1, 28) if ws.cell(row=r, column=c).value is not None]
    print("  r%d: %s" % (r, " | ".join(vals)))

print("\n== COLUMNAS DE DINERO: valores no vacios ==")
for c in range(21, 28):
    n = sum(1 for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=c).value is not None)
    ejemplos = [ws.cell(row=r, column=c).value for r in range(2, ws.max_row + 1)
                if ws.cell(row=r, column=c).value is not None][:3]
    print(f"  {ws.cell(row=1,column=c).value!r}: {n} celdas | ej {ejemplos}")
