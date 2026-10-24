"""Verifica si las formulas TOTAL/venta referencian filas de la misma fecha."""
import re
import sys
from datetime import date, timedelta
from collections import Counter
from openpyxl import load_workbook

path = sys.argv[1]
wbf = load_workbook(path, data_only=False)
ws = wbf["prod1"]

def a_fecha(v):
    if isinstance(v, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(v))
    if isinstance(v, str) and "/" in v:
        try:
            d, m, y = v.strip().split("/")
            y = int(y) + (2000 if len(y) == 2 else 0)
            return date(y, int(m), int(d))
        except Exception:
            return None
    return None

def filas_referenciadas(formula):
    """Extrae filas de un formula tipo =SUM(C106:C111)+C97+C100."""
    filas = set()
    for m in re.finditer(r"([A-Z]{1,2})(\d+):([A-Z]{1,2})(\d+)", formula):
        a, b = int(m.group(2)), int(m.group(4))
        for r in range(min(a, b), max(a, b) + 1):
            filas.add(r)
    # referencias sueltas
    for m in re.finditer(r"(?<![A-Z0-9:])([A-Z]{1,2})(\d+)(?![\d:])", formula):
        filas.add(int(m.group(2)))
    return filas

total_refs = Counter()
malas = []
for r in range(2, ws.max_row + 1):
    b = str(ws.cell(row=r, column=2).value).strip().upper()
    if b not in ("TOTAL", "VENTA"):
        continue
    f = ws.cell(row=r, column=3).value
    if not isinstance(f, str) or not f.startswith("="):
        continue
    fecha_fila = a_fecha(ws.cell(row=r, column=1).value)
    if fecha_fila is None:
        continue
    refs = filas_referenciadas(f)
    otras = []
    for rr in sorted(refs):
        if rr == r:
            continue
        fd = a_fecha(ws.cell(row=rr, column=1).value)
        if fd is not None and fd != fecha_fila:
            otras.append((rr, fd))
    total_refs[b] += 1
    if otras:
        malas.append((r, b, fecha_fila, f[:60], otras[:4]))

print(f"Filas TOTAL con formula: {total_refs['TOTAL']}")
print(f"Filas VENTA con formula: {total_refs['VENTA']}")
print(f"\nFilas que referencian OTRA fecha: {len(malas)}")
for m in malas[:12]:
    print(f"  r{m[0]} {m[1]} fecha={m[2]} formula={m[3]!r}")
    for rr, fd in m[4]:
        print(f"      -> referencia r{rr} que es {fd}")

# formulas que no son SUM puro en TOTAL
raras = []
for r in range(2, ws.max_row + 1):
    if str(ws.cell(row=r, column=2).value).strip().upper() != "TOTAL":
        continue
    f = ws.cell(row=r, column=3).value
    if isinstance(f, str) and not re.fullmatch(r"=SUM\([A-Z]+\d+:[A-Z]+\d+\)", f):
        raras.append((r, f))
print(f"\nTOTAL con formula NO estandar (SUM de rango simple): {len(raras)}")
for r, f in raras[:15]:
    print(f"  r{r}: {f}")
